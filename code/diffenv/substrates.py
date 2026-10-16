"""Loaders that turn the ant and human deposits into the long table `measures.py` scores.

Rebuilt 30 September 2026. The 13 August script that produced the ant (R = 0.512 / 0.560 /
0.548) and InVS (R = 0.370 / 0.312) numbers was never committed and is on neither the
folder nor the OWC drive; this module is the repair, and its first job is to reproduce
those numbers so that all three substrates are scored by one code path.

Ants — Alciatore et al. 2021 `IndividualBehavioursData.xlsx` (one sheet). One row per ant
per observation window (`TimeStamp`, nine windows over ~51 h). Naive ants only (`Treatment ==
'N'`): the immune-challenged `Z` and sham `C` ants are excluded so treatment cannot drive the
result. Constructs, exactly as the sheet's own formula columns define them:
    activity  = nObsActive / Hours       isolation = nObsAlone / Hours       speed = meanSpeed
org = colony, window = TimeStamp, agent = colony:individual.

Humans — SocioPatterns InVS 2013 and 2015, `t i j` rows, one per 20-second slice in which
i and j were in contact. Timestamps are RELATIVE seconds (not epoch); a day is t // 86400.
Construct: contact slices per person per day. org = the deployment, window = day index among
the days that carry records, agent = person id. A person with no contact on a day is treated
as absent that day (no row), not as zero — the deposit cannot distinguish "present and
isolated" from "not in the building". People are kept if they appear on at least two days.
Department labels come from the metadata file and are returned for the specialization join.

Every choice above is a choice; where the August numbers depended on a different one, the
reproduction test will say so rather than the numbers being quietly matched.
"""
from __future__ import annotations

import gzip
from collections import defaultdict
from typing import Dict, List, Tuple

Row = Tuple[str, int, str, float]


def ant_tables(xlsx_path: str, naive_only: bool = True) -> Dict[str, List[Row]]:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=False)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    ix = {h: i for i, h in enumerate(header)}
    need = ("Colony", "Individual", "TimeStamp", "Treatment", "meanSpeed", "nObsAlone", "nObsActive", "Hours")
    missing = [c for c in need if c not in ix]
    if missing:
        raise ValueError(f"workbook lacks columns {missing}; header is {header}")
    tabs = {"activity": [], "isolation": [], "speed": []}
    counts = defaultdict(int)
    windows = sorted({r[ix["TimeStamp"]] for r in rows[1:] if r[ix["TimeStamp"]] is not None})
    widx = {w: i for i, w in enumerate(windows)}
    for r in rows[1:]:
        if r[ix["Colony"]] is None:
            continue
        treat = r[ix["Treatment"]]
        counts[treat] += 1
        if naive_only and treat != "N":
            continue
        org = f"colony{int(r[ix['Colony']])}"
        agent = f"{org}:{int(r[ix['Individual']])}"
        w = widx[r[ix["TimeStamp"]]]
        hours = float(r[ix["Hours"]])
        if hours <= 0:
            continue
        tabs["activity"].append((org, w, agent, float(r[ix["nObsActive"]]) / hours))
        tabs["isolation"].append((org, w, agent, float(r[ix["nObsAlone"]]) / hours))
        tabs["speed"].append((org, w, agent, float(r[ix["meanSpeed"]])))
    tabs["_meta"] = {"rows_by_treatment": dict(counts), "windows": windows,
                     "n_agents": len({a for _, _, a, _ in tabs["activity"]}),
                     "n_colonies": len({o for o, _, _, _ in tabs["activity"]})}
    return tabs


TREATED_IDS = ("BB", "RR")   # BB = immune-challenged (Z), RR = sham (C) in every colony;
                             # verified against the workbook's Treatment column, 30 Sep 2026.


def ant_interaction_tables(network_dir: str, treated=TREATED_IDS) -> Dict[str, object]:
    """Ant `interaction` from the Alciatore `Network_1mmDist6h` matrices — the construct that
    matches the human contact count. One folder per colony, one 9 x 9 symmetric matrix per
    window; a cell is the total time (seconds) two ants spent within 1 mm in that window.

    Per naive ant per window, two versions of "how much did this ant interact":
      interaction        row sum over ALL nestmates (incl. the challenged and sham ants) — the
                         analogue of a person's total contact volume;
      interaction_naive  row sum over NAIVE nestmates only — removes any pull exerted by the
                         treated ants, so treatment cannot drive the result.
    Treated ants are never scored as agents. org = colony folder, window = hour stamp (H_xx),
    agent = colony:ID. Windows are 3 h (H_06, H_51) or 6 h; values are left in seconds because
    the protocol standardizes within colony x window.
    """
    import glob, os, re
    tabs = {"interaction": [], "interaction_naive": []}
    files = sorted(glob.glob(os.path.join(network_dir, "*", "*_InteractionTabNewTmpWin.txt")))
    if not files:
        raise ValueError(f"no network matrices under {network_dir}")
    hours = sorted({int(re.search(r"_H_(\d+)_", os.path.basename(f)).group(1)) for f in files})
    widx = {h: i for i, h in enumerate(hours)}
    per_colony = defaultdict(int)
    for f in files:
        org = os.path.basename(os.path.dirname(f))
        w = widx[int(re.search(r"_H_(\d+)_", os.path.basename(f)).group(1))]
        lines = [l.strip() for l in open(f) if l.strip()]
        ids = lines[0].split(",")[1:]
        m = {}
        for l in lines[1:]:
            p = l.split(",")
            if len(p) != len(ids) + 1:
                raise ValueError(f"ragged matrix in {f}")
            m[p[0]] = dict(zip(ids, (float(x) for x in p[1:])))
        if sorted(m) != sorted(ids):
            raise ValueError(f"row and column ids differ in {f}")
        for a in ids:
            for b in ids:
                if abs(m[a][b] - m[b][a]) > 1e-6:
                    raise ValueError(f"asymmetric matrix in {f}")
        naive = [a for a in ids if a not in treated]
        per_colony[org] += 1
        for a in naive:
            tabs["interaction"].append((org, w, f"{org}:{a}", sum(m[a][b] for b in ids if b != a)))
            tabs["interaction_naive"].append((org, w, f"{org}:{a}", sum(m[a][b] for b in naive if b != a)))
    tabs["_meta"] = {"files": len(files), "colonies": len(per_colony), "window_hours": hours,
                     "windows_per_colony": dict(per_colony), "treated_excluded": list(treated),
                     "n_agents": len({a for _, _, a, _ in tabs["interaction"]})}
    return tabs


def _open(path):
    return gzip.open(path, "rt") if path.endswith(".gz") else open(path)


def invs_tables(tij_path: str, meta_path: str, org: str, min_days: int = 2) -> Dict[str, object]:
    slices = defaultdict(int)          # (person, day) -> contact slices
    days_seen = set(); people = set(); n_rows = 0
    with _open(tij_path) as f:
        for line in f:
            p = line.split()
            if len(p) < 3:
                continue
            t, i, j = int(p[0]), p[1], p[2]
            d = t // 86400
            slices[(i, d)] += 1; slices[(j, d)] += 1
            days_seen.add(d); people.add(i); people.add(j); n_rows += 1
    day_index = {d: k for k, d in enumerate(sorted(days_seen))}
    dept = {}
    with open(meta_path) as f:
        for line in f:
            p = line.split()
            if len(p) >= 2:
                dept[p[0]] = p[1]
    per_person_days = defaultdict(int)
    for (i, d) in slices:
        per_person_days[i] += 1
    keep = {i for i, n in per_person_days.items() if n >= min_days}
    table: List[Row] = [(org, day_index[d], i, float(v)) for (i, d), v in slices.items() if i in keep]
    return {"interaction": table, "department": dept,
            "_meta": {"contact_rows": n_rows, "people_in_contacts": len(people), "people_in_metadata": len(dept),
                      "people_kept": len(keep), "days": len(days_seen), "min_days": min_days,
                      "span_days": (max(days_seen) - min(days_seen) + 1)}}
